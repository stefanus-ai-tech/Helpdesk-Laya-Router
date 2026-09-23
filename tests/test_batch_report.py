import os
import unittest
from io import BytesIO
from pathlib import Path
from threading import Event
from time import sleep
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.batch import BatchInputError, classify_rows, parse_tickets_csv
from backend.reporting.workbook import workbook_bytes
from backend.api.main import app
from scripts.route_github_json import sample_issues


class BatchReportTests(unittest.TestCase):
    def test_sample_csv_and_duplicate_rejection(self):
        sample = Path("dataset/tickets.csv").read_bytes()
        rows = parse_tickets_csv(sample)
        self.assertEqual(len(rows), 100)
        self.assertEqual(rows[0]["ticket_id"], "TCK-0001")
        with self.assertRaisesRegex(BatchInputError, "Duplicate ticket ID"):
            parse_tickets_csv(b"ticket_id,subject,body\nA,Hello,Example message\nA,Again,Another message\n")

    def test_report_has_typed_values_and_queue_sheets(self):
        class Classifier:
            def predict(self, ticket):
                return {"department": "Billing", "department_confidence": .94, "intent": "Refund Request",
                        "intent_confidence": .9, "urgency": "Medium", "urgency_confidence": .8,
                        "frustration": 1, "frustration_confidence": .8, "refund_probability": .96,
                        "churn_probability": .12, "escalation_probability": .1, "source": "laya", "model": "test"}

        rows = parse_tickets_csv(b"ticket_id,subject,body,department,intent,urgency,refund,churn,escalation\n"
                                 b"A,Refund,Please refund this invoice,Billing,Refund Request,Medium,true,false,false\n"
                                 b"B,Another invoice,Please check this invoice,Billing,Payment Problem,Medium,false,false,false\n")
        results = classify_rows(rows, Classifier())
        workbook = load_workbook(BytesIO(workbook_bytes(results, "test.csv")))
        self.assertEqual(workbook.sheetnames, ["Overview", "All Tickets", "Human Review", "Auto Routed", "Evaluation"])
        self.assertEqual(workbook["All Tickets"].max_row, 7)
        self.assertEqual(workbook["All Tickets"]["G6"].value, .94)
        self.assertEqual(workbook["Overview"]["A8"].value, 2)
        self.assertEqual(workbook["Evaluation"]["F8"].value, 2)
        self.assertEqual(workbook["All Tickets"].freeze_panes, "C6")
        self.assertEqual(len(workbook["Overview"]._charts), 1)

    def test_csv_formula_is_escaped(self):
        class Classifier:
            def predict(self, ticket):
                return {"department": "Billing", "department_confidence": .94, "intent": "Other",
                        "intent_confidence": .8, "urgency": "Low", "urgency_confidence": .8,
                        "frustration": 0, "frustration_confidence": .8, "refund_probability": .1,
                        "churn_probability": .1, "escalation_probability": .1, "source": "laya", "model": "test"}

        rows = [{"ticket_id": "A", "subject": "=1+1", "body": "@SUM(1,1) is in this ticket"}]
        workbook = load_workbook(BytesIO(workbook_bytes(classify_rows(rows, Classifier()))))
        self.assertEqual(workbook["All Tickets"]["B6"].data_type, "s")
        self.assertEqual(workbook["All Tickets"]["S6"].data_type, "s")

    def test_upload_returns_excel_in_cuda_mode(self):
        class Classifier:
            def predict(self, ticket):
                return {"department": "Billing", "department_confidence": .94, "intent": "Other",
                        "intent_confidence": .8, "urgency": "Low", "urgency_confidence": .8,
                        "frustration": 0, "frustration_confidence": .8, "refund_probability": .1,
                        "churn_probability": .1, "escalation_probability": .1, "source": "laya", "model": "test"}

        content = b"ticket_id,subject,body\nA,Invoice question,Please explain this invoice\n"
        with patch.dict(os.environ, {"LAYADESK_MODE": "laya"}), patch("torch.cuda.is_available", return_value=True), patch("backend.api.main.classifier", Classifier()):
            with TestClient(app) as client:
                response = client.post("/api/batches/classify", content=content, headers={"Content-Type": "text/csv"})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.headers["content-type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        workbook = load_workbook(BytesIO(response.content))
        self.assertEqual(workbook["All Tickets"]["A6"].value, "A")

    def test_github_sample_caps_repositories_and_preserves_source(self):
        issues = [{"id": i, "title": f"Issue {i}", "body": "A real issue description with enough detail to classify the problem. " * 2,
                   "repo_name": "org/one" if i < 5 else "org/two", "html_url": f"https://github.com/org/repo/issues/{i}",
                   "user": {"login": f"user{i}"}, "labels": []} for i in range(8)]
        rows = sample_issues(issues, 4, 42)
        self.assertEqual(len(rows), 4)
        self.assertEqual(len({row["ticket_id"] for row in rows}), 4)
        self.assertEqual(sum(row["repo_name"] == "org/one" for row in rows), 2)
        self.assertTrue(all(row["issue_url"].startswith("https://github.com/") for row in rows))

    def test_batch_job_reports_real_per_ticket_progress(self):
        second_started = Event()
        release_second = Event()

        class Classifier:
            calls = 0

            def predict(self, ticket):
                self.calls += 1
                if self.calls == 2:
                    second_started.set()
                    release_second.wait(3)
                return {"department": "Billing", "department_confidence": .94, "intent": "Other",
                        "intent_confidence": .8, "urgency": "Low", "urgency_confidence": .8,
                        "frustration": 0, "frustration_confidence": .8, "refund_probability": .1,
                        "churn_probability": .1, "escalation_probability": .1, "source": "laya", "model": "test"}

        content = b"ticket_id,subject,body\nA,Invoice question,Please explain this invoice\nB,Another invoice,Please explain the other invoice\n"
        with patch.dict(os.environ, {"LAYADESK_MODE": "laya"}), patch("torch.cuda.is_available", return_value=True), patch("backend.api.main.classifier", Classifier()):
            with TestClient(app) as client:
                try:
                    started = client.post("/api/batches/jobs", content=content, headers={"Content-Type": "text/csv"})
                    self.assertEqual(started.status_code, 202, started.text)
                    job_id = started.json()["job_id"]
                    self.assertEqual(started.json()["total"], 2)
                    self.assertTrue(second_started.wait(3))
                    running = client.get(f"/api/batches/jobs/{job_id}").json()
                    self.assertEqual(running["processed"], 1)
                    self.assertEqual(running["status"], "running")
                finally:
                    release_second.set()
                for _ in range(100):
                    finished = client.get(f"/api/batches/jobs/{job_id}").json()
                    if finished["status"] == "complete":
                        break
                    sleep(.02)
                self.assertEqual(finished["status"], "complete", finished)
                self.assertEqual(finished["processed"], 2)
                report = client.get(f"/api/batches/jobs/{job_id}/download")
                self.assertEqual(report.status_code, 200)
                workbook = load_workbook(BytesIO(report.content))
                self.assertEqual(workbook["Overview"]["A8"].value, 2)
