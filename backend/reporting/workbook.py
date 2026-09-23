"""Python-generated, reader-friendly Excel report for a Laya CSV batch."""
from __future__ import annotations

from collections import Counter
from datetime import datetime, timezone
from io import BytesIO
from urllib.parse import urlparse

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from backend.evaluation.metrics import calculate
from backend.routing.engine import DEPARTMENTS

INK = "29213D"
MUTED = "7F7790"
PURPLE = "6F42C8"
PALE = "F4EFFC"
LAVENDER = "E8DDF8"
LINE = "E9E5F0"
WHITE = "FFFFFF"
MINT = "E7F7EF"
GREEN = "27845A"
AMBER = "FFF2DF"
ORANGE = "A76722"
RED = "B74C57"
BLUE = "EDF3FF"
FONT = "Aptos"

STATUS_FILL = {"AUTO_ROUTED": MINT, "NEEDS_REVIEW": AMBER}
STATUS_FONT = {"AUTO_ROUTED": GREEN, "NEEDS_REVIEW": ORANGE}
GATE_TEXT = {
    "auto_route": "High confidence",
    "human_confirmation": "Confirm model route",
    "manual_triage": "Low confidence",
    "critical_escalation": "Model flagged critical",
    "policy_critical": "Incident safety rule",
    "senior_confirmation": "Senior confirmation",
    "policy_escalation": "Cancellation language",
    "refund_conflict": "Refund wording conflict",
}


def _text(value) -> str:
    """Keep imported CSV content as text, never as an Excel formula."""
    value = "" if value is None else str(value)
    if value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value[:32767]


def _date(value: str):
    if not value:
        return None
    try:
        parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
        return parsed.astimezone(timezone.utc).replace(tzinfo=None) if parsed.tzinfo else parsed
    except ValueError:
        return _text(value)


def _base_sheet(ws, title: str, subtitle: str, last_col: int):
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.tabColor = PURPLE
    ws.row_dimensions[1].height = 12
    ws.row_dimensions[2].height = 31
    ws.row_dimensions[3].height = 23
    ws.row_dimensions[4].height = 12
    ws["A2"] = title
    ws["A2"].font = Font(name=FONT, size=18, bold=True, color=INK)
    ws["A3"] = subtitle
    ws["A3"].font = Font(name=FONT, size=10, italic=True, color=MUTED)
    ws["A4"].border = Border(bottom=Side(style="thin", color=LINE))
    for cell in ws[4][:last_col]:
        cell.border = Border(bottom=Side(style="thin", color=LINE))
    ws.sheet_view.tabSelected = False


def _section(ws, row: int, label: str, last_col: int):
    ws.row_dimensions[row].height = 23
    for cells in ws.iter_rows(min_row=row, max_row=row, min_col=1, max_col=last_col):
        for cell in cells:
            cell.fill = PatternFill("solid", fgColor=PALE)
            cell.border = Border(bottom=Side(style="thin", color=LAVENDER))
    cell = ws.cell(row, 1, label)
    cell.font = Font(name=FONT, size=10, bold=True, color=PURPLE)
    cell.alignment = Alignment(vertical="center")


def _table_header(ws, row: int, headers: list[str]):
    ws.row_dimensions[row].height = 28
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row, col, header)
        cell.fill = PatternFill("solid", fgColor=INK)
        cell.font = Font(name=FONT, size=9, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(right=Side(style="thin", color="625875"))


def _data_cell(cell, value, row: int, *, numeric=False, percent=False, wrap=False):
    cell.value = value if numeric or isinstance(value, datetime) else _text(value)
    cell.fill = PatternFill("solid", fgColor=WHITE if row % 2 else "FBFAFD")
    cell.font = Font(name=FONT, size=9, color=INK)
    cell.alignment = Alignment(vertical="top", horizontal="right" if numeric else "left", wrap_text=wrap)
    cell.border = Border(bottom=Side(style="hair", color=LINE))
    if percent:
        cell.number_format = "0.0%"
    if isinstance(value, datetime):
        cell.number_format = "yyyy-mm-dd hh:mm"


def _style_status(cell, value: str):
    if value in STATUS_FILL:
        cell.fill = PatternFill("solid", fgColor=STATUS_FILL[value])
        cell.font = Font(name=FONT, size=9, bold=True, color=STATUS_FONT[value])


def _style_priority(cell, value: str):
    if value == "P1":
        cell.fill = PatternFill("solid", fgColor="FCEAED")
        cell.font = Font(name=FONT, size=9, bold=True, color=RED)
    elif value == "P2":
        cell.fill = PatternFill("solid", fgColor=AMBER)
        cell.font = Font(name=FONT, size=9, bold=True, color=ORANGE)


def _source_link(cell, value: str):
    parsed = urlparse(value)
    if parsed.scheme == "https" and parsed.netloc == "github.com":
        cell.hyperlink = value
        cell.font = Font(name=FONT, size=9, color=PURPLE, underline="single")


def _add_filter_and_panes(ws, header_row: int, last_col: int, last_row: int, freeze: str):
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(last_col)}{max(header_row + 1, last_row)}"
    ws.freeze_panes = freeze
    ws.print_title_rows = f"1:{header_row}"


def _all_tickets(wb: Workbook, results: list[dict]):
    ws = wb.create_sheet("All Tickets")
    headers = ["Ticket ID", "Subject", "Status", "Priority", "Queue", "Department", "Dept confidence",
               "Intent", "Intent confidence", "Urgency", "Frustration", "Refund P", "Churn P",
               "Escalation P", "Customer tier", "Customer ID", "Channel", "Created UTC", "Message", "Tags", "Model",
               "Repository", "Issue URL", "GitHub labels", "Issue state"]
    _base_sheet(ws, "All tickets", "Laya CUDA batch results · one row per incoming ticket", len(headers))
    _table_header(ws, 5, headers)
    widths = [15, 34, 18, 11, 31, 22, 17, 22, 17, 14, 13, 13, 13, 15, 17, 17, 13, 20, 67, 33, 16, 31, 68, 35, 15]
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    for index, item in enumerate(results, 6):
        source, p, route = item["input"], item["prediction"], item["route"]
        values = [source["ticket_id"], source["subject"], route["status"], route["priority"], route["queue"],
                  p["department"], p["department_confidence"], p["intent"], p["intent_confidence"],
                  p["urgency"], p["frustration"], p["refund_probability"], p["churn_probability"],
                  p["escalation_probability"], source.get("customer_tier") or "Standard", source.get("customer_id", ""),
                  source.get("source_channel", ""), _date(source.get("created_at", "")), source["body"],
                  ", ".join(route["tags"]), p["model"], source.get("repo_name", ""),
                  source.get("issue_url", ""), source.get("github_labels", ""), source.get("issue_state", "")]
        ws.row_dimensions[index].height = 35
        for col, value in enumerate(values, 1):
            _data_cell(ws.cell(index, col), value, index,
                       numeric=col in (7, 9, 11, 12, 13, 14), percent=col in (7, 9, 12, 13, 14), wrap=col in (2, 19))
        _style_status(ws.cell(index, 3), route["status"])
        _style_priority(ws.cell(index, 4), route["priority"])
        _source_link(ws.cell(index, 23), source.get("issue_url", ""))
        ws.cell(index, 1).font = Font(name=FONT, size=9, bold=True, color=PURPLE)
    last = len(results) + 5
    _add_filter_and_panes(ws, 5, len(headers), last, "C6")
    ws.conditional_formatting.add(f"G6:G{last}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="A782E8", showValue=True))
    return ws


def _queue_sheet(wb: Workbook, name: str, results: list[dict], subtitle: str):
    ws = wb.create_sheet(name)
    headers = ["Ticket ID", "Subject", "Priority", "Department", "Confidence", "Queue", "Reason",
               "Refund P", "Churn P", "Escalation P", "Customer tier", "Message", "Repository", "Issue URL"]
    _base_sheet(ws, name, subtitle, len(headers))
    _table_header(ws, 5, headers)
    widths = [15, 35, 12, 23, 16, 32, 27, 13, 13, 16, 17, 69, 31, 68]
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    for index, item in enumerate(results, 6):
        source, p, route = item["input"], item["prediction"], item["route"]
        values = [source["ticket_id"], source["subject"], route["priority"], p["department"],
                  p["department_confidence"], route["queue"], GATE_TEXT.get(route["gate"], route["gate"]),
                  p["refund_probability"], p["churn_probability"], p["escalation_probability"],
                  source.get("customer_tier") or "Standard", source["body"], source.get("repo_name", ""),
                  source.get("issue_url", "")]
        ws.row_dimensions[index].height = 35
        for col, value in enumerate(values, 1):
            _data_cell(ws.cell(index, col), value, index,
                       numeric=col in (5, 8, 9, 10), percent=col in (5, 8, 9, 10), wrap=col in (2, 12))
        _style_priority(ws.cell(index, 3), route["priority"])
        _source_link(ws.cell(index, 14), source.get("issue_url", ""))
        ws.cell(index, 1).font = Font(name=FONT, size=9, bold=True, color=PURPLE)
    if not results:
        ws["A6"] = "No tickets in this queue."
        ws["A6"].font = Font(name=FONT, size=10, color=MUTED)
    last = len(results) + 5
    _add_filter_and_panes(ws, 5, len(headers), last, "C6")
    if results:
        ws.conditional_formatting.add(f"E6:E{last}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="A782E8", showValue=True))
    ws.sheet_properties.tabColor = "D9994D" if name == "Human Review" else "55AD7D"
    return ws


def _evaluation_rows(results: list[dict]) -> list[dict]:
    labeled = []
    for item in results:
        source = item["input"]
        if not all(source.get(key, "") for key in ("department", "intent", "urgency", "refund", "churn", "escalation")):
            continue
        p, route = item["prediction"], item["route"]
        labeled.append({
            "status": route["status"], "priority": route["priority"],
            "pred_department": p["department"], "pred_intent": p["intent"], "pred_urgency": p["urgency"],
            "refund_probability": p["refund_probability"], "churn_probability": p["churn_probability"],
            "escalation_probability": p["escalation_probability"],
            "department_confidence": p["department_confidence"], "source": "laya",
            "actual_department": source["department"], "actual_intent": source["intent"],
            "actual_urgency": source["urgency"], "actual_refund": source["refund"].lower() == "true",
            "actual_churn": source["churn"].lower() == "true", "actual_escalation": source["escalation"].lower() == "true",
            "origin": "csv_label",
        })
    return labeled


def _evaluation_sheet(wb: Workbook, results: list[dict], source_name: str):
    ws = wb.create_sheet("Evaluation")
    _base_sheet(ws, "Evaluation", f"Ground-truth comparison for {source_name}, when labels are available", 9)
    for col in "ABCDEFGHI":
        ws.column_dimensions[col].width = 18 if col != "A" else 25
    labeled = _evaluation_rows(results)
    if not labeled:
        ws["A6"] = "No complete ground-truth labels were supplied in the CSV."
        ws["A6"].font = Font(name=FONT, size=11, color=MUTED)
        return ws, {"evaluated": 0}
    metrics = calculate(labeled)
    _section(ws, 6, "Classification accuracy", 9)
    accuracy = [
        ("Department accuracy", metrics["department_accuracy"]),
        ("Intent accuracy", metrics["intent_accuracy"]),
        ("Urgency accuracy", metrics["urgency_accuracy"]),
        ("Model critical recall", metrics["model_critical_recall"]),
        ("Operational P1 recall", metrics["operational_p1_recall"]),
        ("Wrong auto-route rate", metrics["wrong_auto_route_rate"]),
    ]
    for row, (label, value) in enumerate(accuracy, 8):
        ws.cell(row, 1, label)
        ws.cell(row, 1).font = Font(name=FONT, size=10, color=INK)
        ws.cell(row, 3, value)
        ws.cell(row, 3).font = Font(name=FONT, size=11, bold=True, color=PURPLE)
        ws.cell(row, 3).number_format = "0.0%" if value is not None else "General"
    ws["E8"] = "Labeled tickets"
    ws["F8"] = metrics["evaluated"]
    ws["E9"] = "Auto-routed tickets"
    ws["F9"] = metrics["auto_routed_count"]
    ws["E10"] = "Critical tickets"
    ws["F10"] = metrics["critical_support"]
    ws["E11"] = "Mean dept confidence"
    ws["F11"] = metrics["average_department_confidence"]
    ws["F11"].number_format = "0.0%"
    for row in range(8, 12):
        for col in (5, 6):
            ws.cell(row, col).font = Font(name=FONT, size=10, bold=col == 6, color=PURPLE if col == 6 else INK)
    _section(ws, 16, "Signal detection", 9)
    _table_header(ws, 18, ["Signal", "Precision", "Recall", "F1", "Support"])
    for row, key, label in ((19, "refund", "Refund request"), (20, "churn", "Churn risk")):
        values = [label, metrics[key]["precision"], metrics[key]["recall"], metrics[key]["f1"], metrics[key]["support"]]
        for col, value in enumerate(values, 1):
            _data_cell(ws.cell(row, col), value, row, numeric=col > 1, percent=col in (2, 3, 4))
    _section(ws, 23, "Department confusion matrix", 9)
    names = list(DEPARTMENTS)
    ws["A25"] = "Actual / predicted"
    for col, name in enumerate(names, 2):
        ws.cell(25, col, name)
    for col in range(1, 7):
        cell = ws.cell(25, col)
        cell.fill = PatternFill("solid", fgColor=INK)
        cell.font = Font(name=FONT, size=9, bold=True, color=WHITE)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[25].height = 32
    matrix = metrics["confusion_matrix"]
    for row, actual in enumerate(names, 26):
        ws.cell(row, 1, actual)
        ws.cell(row, 1).font = Font(name=FONT, size=9, bold=True, color=INK)
        for col, predicted in enumerate(names, 2):
            cell = ws.cell(row, col, matrix.get(actual, {}).get(predicted, 0))
            cell.font = Font(name=FONT, size=10, bold=actual == predicted, color=PURPLE if actual == predicted else INK)
            cell.alignment = Alignment(horizontal="center")
            cell.fill = PatternFill("solid", fgColor=PALE if actual == predicted else WHITE)
    ws["A33"] = "Metrics use raw Laya outputs. P1 recall includes deterministic safety rules."
    ws["A33"].font = Font(name=FONT, size=9, italic=True, color=MUTED)
    return ws, metrics


def _overview(wb: Workbook, results: list[dict], source_name: str, metrics: dict):
    ws = wb.active
    ws.title = "Overview"
    _base_sheet(ws, "Batch triage", f"{source_name} · Laya on CUDA · {len(results)} tickets", 12)
    ws.sheet_properties.tabColor = PURPLE
    for col in "ABCDEFGHIJKL":
        ws.column_dimensions[col].width = 15
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["D"].width = 21
    ws.column_dimensions["G"].width = 21
    ws.column_dimensions["J"].width = 21
    counts = Counter(item["route"]["status"] for item in results)
    confidences = [item["prediction"]["department_confidence"] for item in results]
    cards = [("A7:C9", "A7", "A8", "Tickets", len(results)),
             ("D7:F9", "D7", "D8", "Auto-routed", counts["AUTO_ROUTED"]),
             ("G7:I9", "G7", "G8", "Human review", counts["NEEDS_REVIEW"]),
             ("J7:L9", "J7", "J8", "Mean confidence", sum(confidences) / len(confidences))]
    for area, label_cell, value_cell, label, value in cards:
        for row in ws[area]:
            for cell in row:
                cell.fill = PatternFill("solid", fgColor=PALE)
        ws[label_cell] = label
        ws[label_cell].font = Font(name=FONT, size=10, color=MUTED)
        ws[value_cell] = value
        ws[value_cell].font = Font(name=FONT, size=23, bold=True, color=PURPLE)
        ws[value_cell].number_format = "0%" if label == "Mean confidence" else "#,##0"
    ws.row_dimensions[7].height = 26
    ws.row_dimensions[8].height = 37
    _section(ws, 12, "Routing by department", 12)
    _table_header(ws, 14, ["Department", "Tickets"])
    dept_counts = Counter(item["prediction"]["department"] for item in results)
    for row, department in enumerate(DEPARTMENTS, 15):
        _data_cell(ws.cell(row, 1), department, row)
        _data_cell(ws.cell(row, 2), dept_counts[department], row, numeric=True)
    chart = BarChart()
    chart.type = "bar"
    chart.style = 10
    chart.title = "Assigned department"
    chart.y_axis.title = None
    chart.x_axis.title = "Tickets"
    chart.height = 6
    chart.width = 17
    chart.legend = None
    chart.add_data(Reference(ws, min_col=2, min_row=14, max_row=19), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=15, max_row=19))
    chart.dLbls = DataLabelList()
    chart.dLbls.showVal = True
    ws.add_chart(chart, "D14")
    _section(ws, 27, "Attention signals", 12)
    signals = [
        ("P1 incidents", sum(item["route"]["priority"] == "P1" for item in results)),
        ("Refund score ≥80%", sum(item["prediction"]["refund_probability"] >= .8 for item in results)),
        ("Churn score ≥80%", sum(item["prediction"]["churn_probability"] >= .8 for item in results)),
        ("Escalation score ≥85%", sum(item["prediction"]["escalation_probability"] >= .85 for item in results)),
    ]
    for row, (label, value) in enumerate(signals, 29):
        ws.cell(row, 1, label).font = Font(name=FONT, size=10, color=INK)
        ws.cell(row, 3, value).font = Font(name=FONT, size=11, bold=True, color=PURPLE)
    ws["G29"] = "Labeled tickets"
    ws["I29"] = metrics.get("evaluated", 0)
    ws["G30"] = "Department accuracy"
    ws["I30"] = metrics.get("department_accuracy")
    ws["I30"].number_format = "0.0%"
    ws["G31"] = "Model critical recall"
    ws["I31"] = metrics.get("model_critical_recall")
    ws["I31"].number_format = "0.0%"
    ws["G32"] = "Operational P1 recall"
    ws["I32"] = metrics.get("operational_p1_recall")
    ws["I32"].number_format = "0.0%"
    for row in range(29, 33):
        ws.cell(row, 7).font = Font(name=FONT, size=10, color=INK)
        ws.cell(row, 9).font = Font(name=FONT, size=11, bold=True, color=PURPLE)
    ws["A35"] = "Decision policy"
    ws["A35"].font = Font(name=FONT, size=10, bold=True, color=INK)
    ws["A36"] = "High confidence ≥85% auto-routes. Lower confidence, critical incidents, and safety conflicts go to a person."
    ws["A36"].font = Font(name=FONT, size=9, color=MUTED)
    ws["A38"] = "Source"
    ws["B38"] = _text(source_name)
    ws["A39"] = "Model"
    ws["B39"] = "https://github.com/NandhaKishorM/laya"
    ws["A40"] = "Generated UTC"
    ws["B40"] = datetime.now(timezone.utc).replace(tzinfo=None)
    ws["B40"].number_format = "yyyy-mm-dd hh:mm"
    ws["A41"] = "Labels"
    ws["B41"] = ("CSV labels are supplied ground truth; model results are a batch snapshot."
                 if metrics.get("evaluated") else "No ground-truth labels supplied; accuracy is not calculated.")
    for row in range(38, 42):
        ws.cell(row, 1).font = Font(name=FONT, size=9, bold=True, color=MUTED)
        ws.cell(row, 2).font = Font(name=FONT, size=9, color=INK)
    ws.print_options.horizontalCentered = True
    ws.print_area = "A1:L41"
    return ws


def build_workbook(results: list[dict], source_name: str = "tickets.csv") -> Workbook:
    if not results:
        raise ValueError("Cannot create an empty batch workbook")
    wb = Workbook()
    wb.properties.title = "LayaDesk batch triage"
    wb.properties.subject = "Laya CUDA support ticket classification"
    wb.properties.creator = "LayaDesk"
    _all_tickets(wb, results)
    _queue_sheet(wb, "Human Review", [item for item in results if item["route"]["status"] == "NEEDS_REVIEW"],
                 "Tickets needing confirmation, triage, or escalation")
    _queue_sheet(wb, "Auto Routed", [item for item in results if item["route"]["status"] == "AUTO_ROUTED"],
                 "High-confidence tickets assigned to a queue")
    _, metrics = _evaluation_sheet(wb, results, source_name)
    _overview(wb, results, source_name, metrics)
    wb.active = 0
    return wb


def workbook_bytes(results: list[dict], source_name: str = "tickets.csv") -> bytes:
    output = BytesIO()
    build_workbook(results, source_name).save(output)
    return output.getvalue()
